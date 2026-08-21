"""Import the fund/manager list from an Excel sheet and auto-wire its sources.

Column headers are matched loosely, so the sheet can use whatever wording is
natural. Recognised (case/space insensitive, any one of):

  fund      : fund, fund name, product, product name, scheme, strategy
  house     : house, fund house, amc, pms, company, manager company
  managers  : manager, managers, fund manager, fund managers, cio
  category  : category, type, product type
  youtube   : youtube, youtube channel, channel, yt link
  website   : website, site, url, web, insights page
  keywords  : keywords, extra keywords, match terms, aliases

Only the fund name column is compulsory. Everything else is optional, and any
column that is present gets turned into a source automatically:

  youtube URL -> youtube_channel   (free RSS, catches their own uploads)
  website     -> page_watch        (catches newsletters / factsheets)
  always      -> news_search       (Google News on fund + house + managers)
  managers    -> youtube_search    (catches TV and podcast appearances)

Usage:  python import_funds.py "C:\\path\\to\\funds.xlsx" [--sheet Sheet1] [--dry-run]
"""
import argparse
import re
import sys

import db

ALIASES = {
    "fund": ["fund", "fund name", "product", "product name", "scheme",
             "scheme name", "strategy", "strategy name"],
    "house": ["house", "fund house", "amc", "pms", "pms house", "company",
              "amc name", "manager company", "fund house name"],
    "managers": ["manager", "managers", "fund manager", "fund managers",
                 "fund manager name", "cio", "key personnel"],
    "category": ["category", "type", "product type", "asset class"],
    "youtube": ["youtube", "youtube channel", "youtube link", "channel",
                "yt", "yt link", "youtube url"],
    "website": ["website", "site", "url", "web", "website url", "insights",
                "insights page", "newsletter page"],
    "keywords": ["keywords", "extra keywords", "match terms", "aliases",
                 "other names", "also known as"],
}


def _norm(s):
    return re.sub(r"[^a-z0-9 ]", " ", str(s or "").strip().lower()).strip()


def map_headers(header_row):
    """header cell index -> our field name."""
    mapping = {}
    for idx, cell in enumerate(header_row):
        n = _norm(cell)
        if not n:
            continue
        for field, names in ALIASES.items():
            if n in names and field not in mapping.values():
                mapping[idx] = field
                break
    return mapping


def read_rows(path, sheet=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("sheet is empty")

    # Find the header row: the first row within the top 10 that maps a fund column.
    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:10]):
        m = map_headers(row)
        if "fund" in m.values():
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise SystemExit(
            "could not find a fund/product name column in the first 10 rows.\n"
            f"Headers seen: {[c for c in rows[0] if c]}")

    out = []
    for row in rows[header_idx + 1:]:
        rec = {f: "" for f in ALIASES}
        for idx, field in mapping.items():
            if idx < len(row) and row[idx] is not None:
                rec[field] = str(row[idx]).strip()
        if rec["fund"]:
            out.append(rec)
    return out, mapping


def sources_for(rec):
    """Derive the source list for one imported row."""
    fund, house, managers = rec["fund"], rec["house"], rec["managers"]
    mgr_list = [m.strip() for m in re.split(r"[,;/&]| and ", managers) if m.strip()]
    out = []

    if rec["youtube"]:
        out.append(("youtube_channel", rec["youtube"], f"{house or fund} channel"))

    if rec["website"]:
        out.append(("page_watch", rec["website"], f"{house or fund} website"))

    # One news query covering the product, the house and every manager.
    news_terms = [t for t in ([fund, house] + mgr_list) if t]
    # De-duplicate while preserving order.
    seen, uniq = set(), []
    for t in news_terms:
        if t.lower() not in seen:
            seen.add(t.lower())
            uniq.append(t)
    if uniq:
        query = " OR ".join(f'"{t}"' for t in uniq[:6])
        out.append(("news_search", query, "Google News"))

    # Managers turning up on third-party channels.
    for m in mgr_list[:3]:
        out.append(("youtube_search", m, f"{m} on YouTube"))

    return out


def run(path, sheet=None, dry_run=False):
    db.init()
    rows, mapping = read_rows(path, sheet)
    print(f"Detected columns: {sorted(set(mapping.values()))}")
    print(f"Found {len(rows)} fund row(s).\n")

    for rec in rows:
        srcs = sources_for(rec)
        print(f"  {rec['fund']}"
              + (f"  [{rec['house']}]" if rec["house"] else "")
              + (f"  — {rec['managers']}" if rec["managers"] else ""))
        for kind, value, label in srcs:
            print(f"      {kind:16} {value[:80]}")
        if dry_run:
            continue

        fund_id = db.upsert_fund(
            name=rec["fund"], house=rec["house"], managers=rec["managers"],
            category=rec["category"], match_terms=rec["keywords"])
        for kind, value, label in srcs:
            db.add_source(fund_id, kind, value, label)

    print("\nDry run — nothing written." if dry_run
          else f"\nImported {len(rows)} fund(s). Open /admin to review sources.")
    return len(rows)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", help="path to the .xlsx file")
    ap.add_argument("--sheet", default=None, help="sheet name (default: first)")
    ap.add_argument("--dry-run", action="store_true",
                    help="show what would be imported without writing")
    args = ap.parse_args()
    sys.exit(0 if run(args.path, args.sheet, args.dry_run) else 1)
