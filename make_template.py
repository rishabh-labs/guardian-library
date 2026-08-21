"""Generate funds_template.xlsx - the shape the importer expects.

Column headers are matched loosely, so your existing sheet will probably work
as-is. This is just a reference / starting point.

    python make_template.py
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    ("Fund Name", 34, "Required. The product as you refer to it internally."),
    ("Fund House", 30, "AMC / PMS house name."),
    ("Fund Manager", 34, "Comma separated if more than one."),
    ("Category", 12, "PMS / AIF / MF."),
    ("YouTube Channel", 42, "Paste the URL from the browser address bar."),
    ("Website", 42, "The insights / blog / newsletter page, not the homepage."),
    ("Extra Keywords", 30, "Other names the fund is called in the press."),
]

SAMPLE = [
    ["Marcellus Consistent Compounders", "Marcellus Investment Managers",
     "Saurabh Mukherjea, Rakshit Ranjan", "PMS",
     "", "https://marcellus.in/blog/", "Consistent Compounders, CCP"],
    ["Motilal Oswal Value Strategy", "Motilal Oswal Asset Management",
     "Prateek Agrawal", "PMS",
     "https://www.youtube.com/@MotilalOswalAMC", "", ""],
]


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Funds"

    head_fill = PatternFill("solid", fgColor="1C4ED8")
    head_font = Font(color="FFFFFF", bold=True, size=11)

    for col, (name, width, note) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center")
        cell.comment = None
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(SAMPLE, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    # Second sheet documenting each column, so the template explains itself.
    notes = wb.create_sheet("How to fill this in")
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 90
    notes.cell(row=1, column=1, value="Column").font = Font(bold=True)
    notes.cell(row=1, column=2, value="What to put in it").font = Font(bold=True)
    for i, (name, _w, note) in enumerate(HEADERS, start=2):
        notes.cell(row=i, column=1, value=name)
        notes.cell(row=i, column=2, value=note)
    tail = len(HEADERS) + 3
    notes.cell(row=tail, column=1, value="Note").font = Font(bold=True)
    notes.cell(row=tail, column=2, value=(
        "Only Fund Name is compulsory. Leave anything else blank and the "
        "importer just skips that source. Headers are matched loosely, so "
        "'Product', 'Scheme' or 'Strategy' work in place of 'Fund Name', and "
        "'AMC' works in place of 'Fund House'."))

    wb.save(path)
    return path


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "funds_template.xlsx")
    print("wrote", build(out))
