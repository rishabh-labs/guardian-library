"""Import the Active / Inactive / Knowledge Centre workbook.

The sheet is a watchlist of *strategies*, not fund houses: several rows can
belong to one house (two Tata funds, two Ambit funds, two Abakkus funds). That
distinction drives the whole layout:

  videos      tracked per strategy, so each shelf shows what its own managers
              said - Active, Inactive and Knowledge Centre are separate shelves
  newsletters tracked per fund *house*, because a house publishes one factsheet
              whether we hold one strategy or three

The manager column sometimes holds a YouTube channel URL instead of a name
(Abakkus). Both are handled: a URL becomes a channel subscription, a name
becomes a name search.

    python import_lists.py "<path to Lists.xlsx>" --dry-run
    python import_lists.py "<path to Lists.xlsx>"
"""
import argparse
import re
import sys

import db

SHEET_BUCKETS = {
    "active strategies": "active",
    "inactive strategies": "inactive",
    "knowledge centre": "knowledge",
    "knowledge center": "knowledge",
}

# Rows under this heading in the Knowledge Centre sheet are newsletter sources,
# not video subjects.
# Matches "Newsletters" and the common "Newletters" typo alike.
NEWSLETTER_HEADING = re.compile("new" + chr(92) + "s*s?" + chr(92) + "s*letters?", re.I)
HEADER_WORDS = re.compile(r"fund house|fund manager|names for videos", re.I)

URL = re.compile(r"^https?://", re.I)

# Where the sheet gives a channel URL instead of a person, the house's videos
# still need a named manager: without one, every video on that channel passes
# the "is the manager actually in this?" test by default.
CHANNEL_PEOPLE = {
    "abakkus": "Sunil Singhania",
}

# A Knowledge Centre manager belongs to a house, and that house may publish
# newsletters worth keeping even though we hold none of its strategies.
PERSON_HOUSE = {
    "saurabh mukherjea": "Marcellus",
}


def read_workbook(path):
    """Returns (strategies, kc_people, kc_newsletters).

    strategies      [{name, manager, channel, bucket}]
    kc_people       [name, ...]           - managers followed for learning
    kc_newsletters  [(house, url), ...]
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)

    strategies, kc_people, kc_newsletters = [], [], []

    for sheet in wb.sheetnames:
        bucket = SHEET_BUCKETS.get(sheet.strip().lower())
        if not bucket:
            print(f"  ignoring unrecognised sheet: {sheet}")
            continue
        ws = wb[sheet]
        in_newsletters = False

        for row in ws.iter_rows(values_only=True):
            cells = [("" if c is None else str(c).strip()) for c in row]
            left = next((c for c in cells if c), "")
            if not left:
                continue
            right = ""
            seen_left = False
            for c in cells:
                if not c:
                    continue
                if not seen_left:
                    seen_left = True
                    continue
                right = c
                break

            if NEWSLETTER_HEADING.search(left) and not right:
                in_newsletters = True          # everything after is a feed
                continue
            if HEADER_WORDS.search(left) and not URL.match(right or ""):
                continue                        # column header row

            if bucket == "knowledge" and in_newsletters:
                if right:
                    kc_newsletters.append((left, right))
                continue
            if bucket == "knowledge":
                kc_people.append(left)          # a manager to follow
                continue

            is_url = bool(URL.match(right))
            manager = "" if is_url else right
            if is_url:
                key = house_of(left).lower()
                manager = CHANNEL_PEOPLE.get(key, "")
            strategies.append({
                "name": left,
                "manager": manager,
                "channel": right if is_url else "",
                "bucket": bucket,
            })
    return strategies, kc_people, kc_newsletters


# A strategy name implies its house: "Tata Small Cap" -> Tata, "Ambit Micro
# Cap" -> Ambit. Strip the product words off the end to get there.
PRODUCT_WORDS = re.compile(
    r"\b(flexi\s*cap|flexicap|small\s*cap|smallcap|mid\s*cap|midcap|"
    r"micro\s*cap|microcap|large\s*and\s*mid\s*cap|large\s*cap|largecap|"
    r"focused\s*equity|focused|banking\s*fund|banking|healthcare|"
    r"guardian\s*select|guardian|select|growth|equity|fund|pms|mf|"
    r"mutual\s*fund|india|scheme)\b", re.I)


def house_of(strategy_name):
    core = PRODUCT_WORDS.sub(" ", strategy_name)
    core = re.sub(r"\s+", " ", core).strip(" -,")
    return core or strategy_name.strip()


def run(path, dry_run=False):
    db.init()
    strategies, kc_people, kc_newsletters = read_workbook(path)

    print(f"\n{len(strategies)} strateg(ies), {len(kc_people)} knowledge-centre "
          f"manager(s), {len(kc_newsletters)} knowledge-centre newsletter(s)\n")

    houses = {}
    for s in strategies:
        houses.setdefault(house_of(s["name"]), []).append(s)

    for bucket in ("active", "inactive"):
        rows = [s for s in strategies if s["bucket"] == bucket]
        print(f"=== {bucket.upper()} ({len(rows)}) ===")
        for s in rows:
            who = s["manager"] or f"channel {s['channel']}"
            print(f"  {s['name'][:34]:36} {who[:44]:46} house={house_of(s['name'])}")
        print()

    print(f"=== KNOWLEDGE CENTRE — managers ({len(kc_people)}) ===")
    for p in kc_people:
        print(f"  {p}")
    print(f"\n=== KNOWLEDGE CENTRE — newsletters ({len(kc_newsletters)}) ===")
    for h, u in kc_newsletters:
        print(f"  {h[:28]:30} {u[:64]}")

    if dry_run:
        print("\nDry run — nothing written.")
        return strategies, kc_people, kc_newsletters

    for s in strategies:
        fid = db.upsert_fund(
            name=s["name"], house=house_of(s["name"]),
            managers=s["manager"], category="", match_terms="",
            bucket=s["bucket"])
        if s["channel"]:
            db.add_source(fid, "youtube_channel", s["channel"],
                          f"{s['name']} channel")
        if s["manager"]:
            db.add_source(fid, "youtube_search", s["manager"],
                          f"{s['manager']} on YouTube")

    for person in kc_people:
        fid = db.upsert_fund(name=person,
                             house=PERSON_HOUSE.get(person.lower(), person),
                             managers=person, bucket="knowledge")
        db.add_source(fid, "youtube_search", person, f"{person} on YouTube")

    for house, url in kc_newsletters:
        # "SBI Mutual Fund" and the "SBI" of SBI Banking Fund are one house;
        # normalising here keeps the newsletter shelf from listing both.
        key = house_of(house)
        existing = next((f for f in db.list_funds() if f["house"] == key), None)
        fid = (existing["id"] if existing
               else db.upsert_fund(name=house, house=key, bucket="knowledge"))
        db.add_source(fid, "newsletter_page", url, f"{house} newsletter")

    print(f"\nImported. {len(db.list_funds())} tracked entities.")
    return strategies, kc_people, kc_newsletters


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    sys.exit(0 if run(args.path, args.dry_run) else 1)
